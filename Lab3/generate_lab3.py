import subprocess,sys
subprocess.check_call([sys.executable,"-m","pip","install","cairosvg","reportlab","python-docx"],stdout=subprocess.DEVNULL)
import cairosvg
cairosvg.svg2png(url="Lab3/Lab3_Component_Diagram.svg",write_to="Lab3/Lab3_Component_Diagram_Kushagra_Bhandari_PES1UG24CS245.png",output_width=1600)
cairosvg.svg2pdf(url="Lab3/Lab3_Component_Diagram.svg",write_to="Lab3/Lab3_Component_Diagram_Kushagra_Bhandari_PES1UG24CS245.pdf")
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate,Paragraph
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name="BodyX",parent=styles["BodyText"],fontSize=9.5,leading=12,spaceAfter=4,textColor=colors.HexColor("#334155"))); styles.add(ParagraphStyle(name="HeadX",parent=styles["Heading2"],fontSize=11,leading=13,spaceBefore=5,spaceAfter=3,textColor=colors.HexColor("#4f46e5")))
doc=SimpleDocTemplate("Lab3/Lab3_Written_Justification_Kushagra_Bhandari_PES1UG24CS245.pdf",pagesize=A4,leftMargin=48,rightMargin=48,topMargin=42,bottomMargin=42)
story=[Paragraph("Lab 3 – Written Justification",styles["Title"]),Paragraph("Self-Service Coffee Kiosk | PES1UG24CS245",styles["Heading3"])]
sections=[("Architecture Selection","We chose <b>Layered Architecture</b> for the Self-Service Coffee Kiosk System."),("Architectural Choice","The system is divided into Presentation, Business, and Data/Hardware layers. The Presentation layer handles touch-screen interaction; the Business layer handles ordering and payment; and the Data/Hardware layer handles menu/pricing storage and receipt printing."),("Reason 1 – Clear separation of responsibilities","The kiosk has distinct responsibilities: customer interaction, order/payment processing, and menu/pricing plus printer access. Keeping these responsibilities in separate layers makes the component boundaries clear and simplifies maintenance."),("Reason 2 – Suitable for a focused kiosk workflow","The scenario has a small, fixed workflow: three coffee types, two sizes, and credit-card payment only. A layered design keeps this workflow straightforward without introducing the operational complexity of independently deployed services."),("Security Advantage","Payment processing is isolated in the Payment Service instead of being mixed with the touch-screen or database components. Layer boundaries can restrict access so components expose only the operations required by other components."),("Performance Benefit","The kiosk can communicate directly between local components, avoiding unnecessary network hops and distributed-service overhead. Menu and pricing information can also be accessed locally by the kiosk."),("Conclusion","Layered Architecture provides clear separation, simple deployment, and a straightforward communication path for the kiosk scenario.")]
for h,b in sections: story += [Paragraph(h,styles["HeadX"]),Paragraph(b,styles["BodyX"])]
doc.build(story)
from docx import Document
from docx.shared import Pt, Inches
d=Document(); s=d.sections[0]; s.top_margin=s.bottom_margin=Inches(.55); s.left_margin=s.right_margin=Inches(.65)
p=d.add_paragraph(); p.alignment=1; r=p.add_run("Lab 3 – Written Justification"); r.bold=True; r.font.size=Pt(16)
p=d.add_paragraph("Self-Service Coffee Kiosk | PES1UG24CS245"); p.alignment=1; p.runs[0].bold=True
for h,b in sections:
 p=d.add_paragraph(); r=p.add_run(h); r.bold=True; r.font.size=Pt(10.5)
 p=d.add_paragraph(b.replace("<b>","").replace("</b>","")); p.paragraph_format.space_after=Pt(3); p.runs[0].font.size=Pt(9.5)
d.save("Lab3/Lab3_Written_Justification_Kushagra_Bhandari_PES1UG24CS245.docx")
